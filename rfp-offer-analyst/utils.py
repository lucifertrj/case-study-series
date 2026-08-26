import io
import sqlite3

import openpyxl
import pymupdf
from docx import Document

SCAN_THRESHOLD_CHARS_PER_PAGE = 50
DOCX_CHUNK_CHARS = 1500

STANDARD_CRITERIA = [
    {"question": "Delivery within 30 days of order", "weight": 5, "is_knockout": False},
    {"question": "Payment terms Net 60 or better", "weight": 3, "is_knockout": False},
    {"question": "GDPR data processing addendum included", "weight": 4, "is_knockout": True},
    {"question": "Liability cap at or above contract value", "weight": 4, "is_knockout": True},
    {"question": "ISO 27001 certification held", "weight": 3, "is_knockout": False},
    {"question": "Price sheet complete, all line items", "weight": 4, "is_knockout": False},
    {"question": "Subcontractors disclosed", "weight": 2, "is_knockout": False},
    {"question": "Warranty period of 24 months minimum", "weight": 3, "is_knockout": False},
]


def extract_pdf(data: bytes) -> list[tuple[int, str]]:
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        return [(i + 1, page.get_text()) for i, page in enumerate(doc)]
    finally:
        doc.close()


def extract_docx(data: bytes) -> list[tuple[int, str]]:
    doc = Document(io.BytesIO(data))
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if not full_text:
        return [(1, "")]
    chunks = [
        full_text[i : i + DOCX_CHUNK_CHARS]
        for i in range(0, len(full_text), DOCX_CHUNK_CHARS)
    ]
    return [(i + 1, chunk) for i, chunk in enumerate(chunks)]


def extract_xlsx(data: bytes) -> list[tuple[int, str]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    pages = []
    for sheet_idx, sheet in enumerate(wb.worksheets, start=1):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        pages.append((sheet_idx, f"[Sheet: {sheet.title}]\n" + "\n".join(rows)))
    return pages or [(1, "")]


def is_scanned(pages: list[tuple[int, str]]) -> bool:
    if not pages:
        return False
    avg_chars = sum(len(text) for _, text in pages) / len(pages)
    return avg_chars < SCAN_THRESHOLD_CHARS_PER_PAGE


def extract(filename: str, data: bytes) -> tuple[list[tuple[int, str]], bool]:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        pages = extract_pdf(data)
        return pages, is_scanned(pages)
    if ext == "docx":
        return extract_docx(data), False
    if ext in ("xlsx", "xlsm"):
        return extract_xlsx(data), False
    raise ValueError(f"Unsupported file type: {filename}")


def seed_if_empty(conn: sqlite3.Connection) -> None:
    count = conn.execute("SELECT COUNT(*) FROM criteria").fetchone()[0]
    if count > 0:
        return
    conn.executemany(
        "INSERT INTO criteria (question, weight, is_knockout, source) VALUES (?, ?, ?, 'standard')",
        [(c["question"], c["weight"], int(c["is_knockout"])) for c in STANDARD_CRITERIA],
    )
    conn.commit()


def calculate_scores(conn):
    offers = conn.execute("SELECT * FROM offers ORDER BY id").fetchall()
    criteria = conn.execute("SELECT * FROM criteria ORDER BY id").fetchall()

    if not criteria:
        max_possible_score = 0
    else:
        max_possible_score = sum(c["weight"] * 2 for c in criteria)

    results = []
    for offer in offers:
        ev_rows = conn.execute(
            "SELECT * FROM evaluations WHERE offer_id = ?", (offer["id"],)
        ).fetchall()

        evals_by_crit = {e["criterion_id"]: dict(e) for e in ev_rows}

        score = 0
        unclear_count = 0
        is_disqualified = False
        dq_reasons = []

        for c in criteria:
            cid = c["id"]
            ev = evals_by_crit.get(cid)
            verdict = ev["verdict"] if ev else "unclear"

            if verdict == "full":
                score += 2 * c["weight"]
            elif verdict == "partial":
                score += c["weight"]
            elif verdict == "unclear":
                unclear_count += 1
            elif verdict == "none" and c["is_knockout"]:
                is_disqualified = True
                dq_reasons.append(f"knockout: {c['question']}")

        dq_reason = ", ".join(dq_reasons) if dq_reasons else None

        results.append(
            {
                "offer_id": offer["id"],
                "name": offer["name"],
                "supplier": offer["supplier"],
                "price": offer["price"] or "",
                "score": score,
                "max_score": max_possible_score,
                "is_disqualified": is_disqualified,
                "dq_reason": dq_reason,
                "unclear_count": unclear_count,
                "evaluations": evals_by_crit,
            }
        )

    results.sort(key=lambda x: (1 if x["is_disqualified"] else 0, -x["score"], x["name"]))
    return results, max_possible_score
